import json
import os
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

from conftest import REPO_ROOT


LEGACY_RUNNER = Path(REPO_ROOT) / "factory_core" / "adapters" / "legacy_runner.sh"


def write_file(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


def write_zip(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("paper.pdf", "fake")


def make_project_with_delivery(factory: Path, base: str, verdict: str = "PASS") -> Path:
    project = factory / "ongoing" / base
    paper = project / f"{base}_paper.tex"
    write_file(project / "problem" / "problem_brief.md", "# brief\n")
    write_file(project / "checkpoint.md", "- **Last completed step**: 15\n")
    write_file(project / "judge_evaluation.md", f"VERDICT: {verdict}\n" + "\n".join(["judge"] * 30) + "\n")
    write_file(project / "reviewer_entry_map.md", "# map\n")
    write_file(project / "anchor_figure_plan.md", "# anchors\n")
    write_file(project / "entry_gate.md", "VERDICT: PASS\n")
    write_file(paper, "\\begin{document}\n" + "\n".join(["paper"] * 220) + "\n\\end{document}\n")
    write_file(factory / "papers" / f"{base}_paper.pdf", "%PDF fake\n")
    write_zip(factory / "papers" / f"{base}_submission.zip")
    return project


def test_paper_reviewer_scope_excludes_unseen_pdf_visual_quality():
    prompt = (Path(REPO_ROOT) / "prompts/judges/paper_reviewer.txt").read_text(
        encoding="utf-8"
    )
    human_rubric = (Path(REPO_ROOT) / "evaluation/human_rubric.md").read_text(
        encoding="utf-8"
    )

    assert "不向你提供 PDF 渲染画面或图片像素" in prompt
    assert "不得评价分页、字体" in prompt
    assert "PDF 字节指纹只证明交付版本一致" in prompt
    assert "渲染与视觉复核（独立记录，不进入当前自动六维分数）" in human_rubric


def test_step14_prompt_does_not_fabricate_gate2_pass():
    prompt = (Path(REPO_ROOT) / "prompts/step14_abstract.txt").read_text(
        encoding="utf-8"
    )

    assert "Step 13 Gate 2 已经 PASS" not in prompt
    assert "不得假定 Gate 2 PASS" in prompt
    assert "gate2_delivery_override.json" in prompt
    assert "`gate2_delivery_override.json` 不是技术续跑授权凭据" in prompt
    assert "GATE2_CONTINUATION_AUTHORIZED" in prompt


def test_evaluator_rejects_incomplete_canonical_results(tmp_path, monkeypatch):
    project = tmp_path / "ongoing" / "demo"
    from test_evaluate_modeling_project_step8_5 import make_complete_project

    make_complete_project(project)
    write_file(project / "reviewer_entry_map.md", "# map\n")
    write_file(project / "anchor_figure_plan.md", "# anchors\n")
    write_file(project / "entry_gate.md", "VERDICT: PASS\n")
    write_file(project / "results" / "p1" / "values.json", '{"problem": 1, "status": "CONVERGED", "objective": 1.0}\n')
    write_file(project / "results" / "p2" / "values.json", '{"problem": 2, "status": "RUNNING"}\n')

    from scripts import evaluate_modeling_project as mod

    monkeypatch.setattr(mod, "infer_step", lambda root, project: (16, "16"))
    monkeypatch.setattr(mod, "zip_ok", lambda path: (True, "ok"))

    ev = mod.evaluate(project, tmp_path)
    checks = {check.name: check for check in ev.checks}

    assert checks["canonical_results"].ok is False
    assert "p2" in checks["canonical_results"].detail


def test_evaluator_and_delivery_contract_reject_failed_reality_gates(tmp_path, monkeypatch):
    project = tmp_path / "ongoing" / "demo_reality_gate"
    from test_evaluate_modeling_project_step8_5 import make_complete_project

    make_complete_project(project)
    write_file(tmp_path / "method_library" / "demo.md", "# demo method\n")
    write_file(project / "reviewer_entry_map.md", "# map\n")
    write_file(project / "anchor_figure_plan.md", "# anchors\n")
    write_file(project / "entry_gate.md", "VERDICT: PASS\n")
    write_file(project / "numbers_manifest.json", '{"sources": {}}\n')
    write_file(project / "results" / "canonical_results.json", '{"objective": 1.0, "status": "FEASIBLE"}\n')

    from scripts import delivery_contract
    from scripts import evaluate_modeling_project as mod

    monkeypatch.setattr(mod, "infer_step", lambda root, project: (16, "16"))
    monkeypatch.setattr(mod, "zip_ok", lambda path: (True, "ok"))
    monkeypatch.setattr(mod, "symbol_check_ok", lambda root, project, base: (True, "verify_symbols PASS"))

    def fake_run_python_check(root, args, timeout=60):
        script = args[0]
        if script == "scripts/verify_provenance.py":
            return False, "VERDICT: REPAIR_FALLBACK"
        if script == "scripts/verify_spec_impl.py":
            return False, "VERDICT: FAIL"
        return True, "VERDICT: PASS"

    monkeypatch.setattr(mod, "run_python_check", fake_run_python_check)

    ev = mod.evaluate(project, tmp_path)
    checks = {check.name: check for check in ev.checks}

    assert checks["provenance_gate"].ok is False
    assert checks["spec_impl_gate"].ok is False
    assert ev.passed is False

    manifest = delivery_contract.build_delivery_manifest(project, tmp_path, ev, generated_at="2026-07-10T00:00:00+00:00")

    assert manifest["status"] != "CURRENT_PASS"
    assert {check["name"] for check in manifest["evaluation"]["failed_checks"]} >= {
        "provenance_gate",
        "spec_impl_gate",
    }


def test_evaluator_rejects_failed_project_quality_contract(tmp_path, monkeypatch):
    project = tmp_path / "ongoing" / "demo_quality_contract"
    from test_evaluate_modeling_project_step8_5 import make_complete_project

    make_complete_project(project)
    write_file(tmp_path / "method_library" / "demo.md", "# demo method\n")
    write_file(project / "reviewer_entry_map.md", "# map\n")
    write_file(project / "anchor_figure_plan.md", "# anchors\n")
    write_file(project / "entry_gate.md", "VERDICT: PASS\n")
    write_file(project / "numbers_manifest.json", '{"sources": {}}\n')
    write_file(project / "results" / "canonical_results.json", '{"objective": 1.0, "status": "FEASIBLE"}\n')
    write_file(project / "quality_contract.json", '{"version": 1, "claims": [], "anomaly_checks": []}\n')

    from scripts import evaluate_modeling_project as mod

    monkeypatch.setattr(mod, "infer_step", lambda root, project: (16, "16"))
    monkeypatch.setattr(mod, "zip_ok", lambda path: (True, "ok"))
    monkeypatch.setattr(mod, "symbol_check_ok", lambda root, project, base: (True, "verify_symbols PASS"))

    def fake_run_python_check(root, args, timeout=60):
        if args[0] == "scripts/verify_quality_contract.py":
            return False, "QUALITY_CONTRACT_FAILURES=1\nVERDICT: FAIL"
        return True, "VERDICT: PASS"

    monkeypatch.setattr(mod, "run_python_check", fake_run_python_check)

    ev = mod.evaluate(project, tmp_path)
    checks = {check.name: check for check in ev.checks}

    assert checks["quality_contract_gate"].ok is False
    assert ev.passed is False


def test_verify_numbers_manifest_includes_nested_json_and_xlsx(tmp_path):
    project = tmp_path / "proj"
    write_file(project / "results" / "p1" / "values.json", json.dumps({
        "problem": 1,
        "objective": 4.9,
        "decision": {"theta_deg": 8.7, "v_mps": 140.0},
        "intervals": [[1.5, 6.4]],
    }))
    try:
        from openpyxl import Workbook
    except ImportError:
        return
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "duration"
    ws["B1"] = 14.804
    wb.save(project / "result3.xlsx")

    from verify_numbers import scan_results_directory

    manifest = scan_results_directory(project)
    flat_keys = set(manifest["results/p1/values.json"].keys())

    assert "decision.theta_deg" in flat_keys
    assert "intervals[0][1]" in flat_keys
    assert manifest["result3.xlsx"]["Sheet!B1"]["value"] == 14.804


def test_verify_numbers_rejects_unrecorded_simple_derived_number(tmp_path):
    project = tmp_path / "proj"
    base = "proj"
    write_file(project / "results" / "p1" / "values.json", json.dumps({"a": 2.0, "b": 3.0}))
    write_file(project / f"{base}_paper.tex", "\\begin{document}\nThe reported value is 5.0.\n\\end{document}\n")

    from verify_numbers import generate_manifest, verify_paper

    generate_manifest(project)

    assert verify_paper(project, base) is False
    report = (project / "number_verification.md").read_text(encoding="utf-8")
    assert "5.0" in report


def test_verify_numbers_handles_latex_commands_and_exponents(tmp_path):
    project = tmp_path / "proj"
    base = "proj"
    write_file(project / "results" / "p1" / "values.json", json.dumps({"mask_time_s": 1.3624}))
    write_file(
        project / f"{base}_paper.tex",
        "\\begin{document}\n"
        "The verified result is $T_1\\approx1.3624\\,\\mathrm{s}$ with tolerance $<10^{-6}$.\n"
        "\\end{document}\n",
    )
    from verify_numbers import generate_manifest, verify_paper

    generate_manifest(project)

    assert verify_paper(project, base) is True


def test_verify_numbers_handles_ranges_scientific_notation_and_layout(tmp_path):
    project = tmp_path / "proj"
    base = "proj"
    write_file(
        project / "results" / "p1" / "values.json",
        json.dumps(
            {
                "range": [1200.47, 3999.64],
                "tiny": 4.0779e-7,
            }
        ),
    )
    write_file(
        project / f"{base}_paper.tex",
        "\\begin{document}\n"
        "\\begin{longtable}{p{0.22\\textwidth}p{0.49\\textwidth}}\n"
        "Range: 1200.47--3999.64; spreads: $4.08\\times 10^{-7}$ and $4.08\\times10^-7$.\n"
        "\\end{longtable}\n"
        "\\lstinputlisting[firstline=53,lastline=170]{models/02_model.py}\n"
        "\\end{document}\n",
    )
    write_file(project / "models" / "02_model.py", "# fixture\n")

    from verify_numbers import generate_manifest, verify_paper

    generate_manifest(project)

    assert verify_paper(project, base) is True


def test_verify_numbers_ignores_multiline_lstinputlisting_line_selectors(tmp_path):
    project = tmp_path / "proj"
    base = "proj"
    write_file(
        project / "results" / "p1" / "values.json",
        json.dumps({"result": 42.5}),
    )
    write_file(
        project / f"{base}_paper.tex",
        "\\begin{document}\n"
        "Verified result: 42.5.\n"
        "\\lstinputlisting[\n"
        "  language=Python,\n"
        "  firstline=540,\n"
        "  lastline=650\n"
        "]{models/02_model.py}\n"
        "\\end{document}\n",
    )
    write_file(project / "models" / "02_model.py", "# fixture\n")

    from verify_numbers import generate_manifest, verify_paper

    generate_manifest(project)

    assert verify_paper(project, base) is True


def test_verify_numbers_ignores_tikz_layout_coordinates_but_checks_caption(tmp_path):
    project = tmp_path / "proj"
    base = "proj"
    write_file(project / "results" / "p1" / "values.json", json.dumps({"result": 42.5}))
    write_file(
        project / f"{base}_paper.tex",
        "\\begin{document}\n"
        "\\begin{figure}\n"
        "\\begin{tikzpicture}\n"
        "\\draw (1.25,3.75) -- (8.5,9.5);\n"
        "\\end{tikzpicture}\n"
        "\\caption{Verified result: 42.5.}\n"
        "\\end{figure}\n"
        "\\end{document}\n",
    )

    from verify_numbers import generate_manifest, verify_paper

    generate_manifest(project)

    assert verify_paper(project, base) is True


def test_verify_numbers_rejects_key_result_that_disagrees_with_canonical_source(tmp_path):
    project = tmp_path / "proj"
    base = "proj"
    write_file(project / "results" / "p1" / "values.json", json.dumps({"decision": {"d": 42.5}}))
    write_file(
        project / "results" / "key_results.json",
        json.dumps(
            {
                "key_results": [
                    {
                        "label": "headline thickness",
                        "value": 41.5,
                        "canonical_source": "results/p1/values.json::decision.d",
                    }
                ]
            }
        ),
    )
    write_file(project / f"{base}_paper.tex", "\\begin{document}\nResult: 41.5.\n\\end{document}\n")

    from verify_numbers import generate_manifest, verify_paper

    generate_manifest(project)

    assert verify_paper(project, base) is False
    report = (project / "number_verification.md").read_text(encoding="utf-8")
    assert "Key-Result Canonical Source Mismatches" in report
    assert "headline thickness" in report


def test_verify_symbols_treats_big_set_operators_as_latex_noise(tmp_path):
    project = tmp_path / "proj"
    base = "proj"
    write_file(project / "symbol_table.md", "| 符号 | 含义 |\n|---|---|\n| $I$ | interval |\n| $j$ | index |\n")
    write_file(
        project / f"{base}_paper.tex",
        "\\begin{document}\n"
        "\\section{符号说明}\n"
        "Registered variables are $I$ and $j$.\n"
        "The aggregate is $\\bigcup_{j=1}^{3} I_j$ and $\\bigcap_{j=1}^{3} I_j$.\n"
        "\\end{document}\n",
    )

    from verify_symbols import collect_symbol_metrics

    metrics = collect_symbol_metrics(project, base)

    assert "\\bigcup" not in metrics["_undefined_list"]
    assert "\\bigcap" not in metrics["_undefined_list"]


def test_verify_symbols_ignores_norm_delimiters_and_equation_labels(tmp_path):
    project = tmp_path / "proj"
    base = "proj"
    write_file(
        project / "symbol_table.md",
        "| 符号 | 含义 |\n|---|---|\n| $x$ | vector |\n| $t$ | time |\n",
    )
    write_file(
        project / f"{base}_paper.tex",
        "\\begin{document}\n"
        "\\section{符号说明}\n"
        "Registered variables are $x$ and $t$.\n"
        "\\begin{equation}\n"
        "\\lVert x\\rVert = t.\n"
        "\\label{eq:active_window}\n"
        "\\end{equation}\n"
        "\\end{document}\n",
    )

    from verify_symbols import collect_symbol_metrics

    metrics = collect_symbol_metrics(project, base)

    assert metrics["_undefined_list"] == []


def test_verify_symbols_ignores_unbraced_roman_unit_letter(tmp_path):
    project = tmp_path / "proj"
    base = "proj"
    write_file(
        project / "symbol_table.md",
        "| 符号 | 含义 |\n|---|---|\n| $T$ | temperature |\n",
    )
    write_file(
        project / f"{base}_paper.tex",
        "\\begin{document}\n"
        "\\section{符号说明}\n"
        "The temperature is $T=240\\,{}^\\circ\\mathrm C$.\n"
        "\\end{document}\n",
    )

    from verify_symbols import collect_symbol_metrics

    metrics = collect_symbol_metrics(project, base)

    assert metrics["_undefined_list"] == []


def test_public_runner_is_a_python_compatibility_launcher():
    text = (Path(REPO_ROOT) / "run_paper.sh").read_text(encoding="utf-8")

    assert "python3 -m factory_core.cli compat" in text
    assert "run_step_1()" not in text
    assert len(text.splitlines()) < 40


def test_step6_precheck_parses_markdown_assumption_table(tmp_path):
    project = tmp_path / "proj"
    write_file(project / "assumption_ledger.md", "\n".join([
        "| id | 陈述 | 来源 | 若违反的影响 | 状态 | 标签 |",
        "|---|---|---|---|---|---|",
        "| A1 | assumption | source | impact | INHERITED | **PROTECTED** |",
        "| A2 | assumption | source | impact | OPEN | CRITICAL |",
    ]))

    from scripts.step6_coverage_precheck import check_assumption_ledger

    ok, msg, counts = check_assumption_ledger(project)

    assert ok is True
    assert counts["INHERITED"] == 1
    assert counts["OPEN"] == 1


def test_project_monitor_once_handles_missing_results_dirs():
    project_name = "monitor_missing_results_fixture"
    project = Path(REPO_ROOT) / "ongoing" / project_name
    monitor_log = Path(REPO_ROOT) / "run_state" / f"{project_name}_monitor.log"

    shutil.rmtree(project, ignore_errors=True)
    monitor_log.unlink(missing_ok=True)
    write_file(project / "checkpoint.md", "- **Last completed step**: 0\n")
    write_file(project / "problem" / "problem_brief.md", "# brief\n")
    write_file(project / "logs" / "runner.log", "runner started\n")

    try:
        out = subprocess.run(
            [os.path.join(REPO_ROOT, "scripts", "project_monitor.sh"), "--once", project_name],
            cwd=REPO_ROOT,
            capture_output=True,
            text=True,
            check=False,
        )

        assert out.returncode == 0, out.stderr
        report = monitor_log.read_text(encoding="utf-8")
        assert "values_json_count=0" in report
        assert "solver_job_meta_count=0" in report
    finally:
        shutil.rmtree(project, ignore_errors=True)
        monitor_log.unlink(missing_ok=True)
